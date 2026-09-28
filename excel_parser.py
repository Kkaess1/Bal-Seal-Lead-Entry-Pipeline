import openpyxl
from pathlib import Path

from ai_column_mapper import map_columns_with_ai


# Maps our internal field names to every column header spelling we might encounter.
COLUMN_ALIASES = {
    # A single "Name" column is common in hand-made lists and badge-scan
    # exports. Used only when a row has no separate first/last name -- see
    # _split_full_name().
    'full_name':       ['name', 'full name', 'contact name', 'attendee name',
                        'contact'],
    'first_name':      ['first name', 'firstname', 'first', 'given name'],
    'last_name':       ['last name', 'lastname', 'last', 'surname', 'family name'],
    'organization':    ['organization', 'company', 'org', 'account', 'account name',
                        'employer', 'company name'],
    'title':           ['title', 'job title', 'position', 'role'],
    # Work phone takes priority; mobile is a fallback (see phone_mobile below)
    'phone':           ['phone', 'work phone', 'contact work phone', 'telephone', 'tel',
                        'phone number', 'work phone number'],
    'phone_mobile':    ['mobile phone number', 'mobile phone', 'cell phone', 'cell', 'mobile'],
    'email':           ['email', 'email address', 'contact email', 'e-mail'],
    'address1':        ['address 1', 'address1', 'street', 'street address', 'addr1', 'addr 1',
                        'work address 1', 'work address1'],
    'address2':        ['address 2', 'address2', 'addr2', 'addr 2',
                        'work address 2', 'work address2'],
    'city':            ['city', 'work city'],
    'state':           ['state/province', 'state', 'province', 'state / province',
                        'work state/province', 'work state'],
    'zip':             ['zip/postal code', 'zip', 'postal code', 'postcode', 'zip code',
                        'work zip/postal code'],
    'country':         ['country', 'work country'],
    'notes_action':    ['action'],
    'notes_body':      ['notes', 'note', 'application notes'],
    'survey_answers':  ['survey answers', 'survey', 'answers'],
    'qualifiers':      ['qualifiers', 'qualifier'],
    'rank':            ['rank', 'hot/warmcold', 'priority', 'lead rank', 'lead interest'],
    'scan_device':     ['scan device', 'device', 'scanned by', 'booth staff first name'],
}

# Column 9 (index 8) in the known format is an unnamed rank column.
UNNAMED_RANK_COL_INDEX = 8


# Every header spelling the alias table knows, for telling a header row from a
# title row above it.
_KNOWN_HEADERS = {alias for aliases in COLUMN_ALIASES.values() for alias in aliases}

# The fields that make a row a lead at all -- a sheet mapping none of them
# yields nothing but skipped rows.
_NAME_OR_COMPANY_FIELDS = ('first_name', 'last_name', 'full_name', 'organization')


def _find_header_row(rows):
    """
    Scan the first 10 rows to find the one that looks like a header.

    Takes the sheet's rows already read into lists and returns a zero-based
    index into them. Rows are read once with iter_rows() rather than cell by
    cell: in read-only mode every ws.cell() call re-scans the sheet XML, which
    made a few hundred rows take minutes.

    A row holding two or more known header spellings wins first. "Mostly
    text" alone picked a title row such as "Show: | MD&M West | Booth 12"
    over the real headers under it, and never found a two-column header
    ("Name | Company") at all -- both common in hand-made lists, and both
    ended in an import that found no name or company column.
    """
    for idx, row in enumerate(rows[:10]):
        known = sum(1 for c in row
                    if isinstance(c, str) and c.strip().lower() in _KNOWN_HEADERS)
        if known >= 2:
            return idx, list(row)
    for idx, row in enumerate(rows[:10]):
        # A header row has mostly string values
        text_cells = sum(1 for c in row if isinstance(c, str) and c.strip())
        if text_cells >= 3:
            return idx, list(row)
    return 0, list(rows[0]) if rows else []


def _read_rows(ws):
    """
    Every row of a sheet as a list, all padded to the widest row.

    Padded because a workbook with no <dimension> record (some exporters
    omit it) comes back ragged: a header row that stops at column 8 then has
    no index 8 at all, the unnamed-rank-column rule in _build_column_map()
    never fires, and every Cold row is imported as a lead.
    """
    if not hasattr(ws, 'iter_rows'):   # a chartsheet has no cells
        return []
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    width = max((len(r) for r in rows), default=0)
    for row in rows:
        row.extend([None] * (width - len(row)))
    return rows


def _maps_name_or_company(col_map):
    return any(field in col_map for field in _NAME_OR_COMPANY_FIELDS)


def _build_column_map(headers):
    """
    Return a dict mapping field_name → zero-based column index.
    Matches case-insensitively and ignores leading/trailing whitespace.
    """
    normalized = {
        i: (str(h).strip().lower() if h else '')
        for i, h in enumerate(headers)
    }

    col_map = {}
    for field, aliases in COLUMN_ALIASES.items():
        for i, header_text in normalized.items():
            if header_text in aliases:
                col_map[field] = i
                break

    # Special case: unnamed column at index 8 that contains Hot/Warm/Cold
    if (
        'rank' not in col_map
        and UNNAMED_RANK_COL_INDEX < len(headers)
        and not normalized[UNNAMED_RANK_COL_INDEX]
    ):
        col_map['rank'] = UNNAMED_RANK_COL_INDEX

    # AI fallback: map any fields that alias-matching didn't cover
    ai_additions = map_columns_with_ai(headers, already_mapped=set(col_map.keys()))
    # A column the alias table already claimed is not the model's to reuse.
    # Asked for every unmapped field, it can offer the Company column as
    # full_name too, and every company-only row would gain a contact named
    # after its employer.
    claimed = set(col_map.values())
    ai_additions = {f: i for f, i in (ai_additions or {}).items() if i not in claimed}
    if ai_additions:
        print(f'  AI mapped    : {", ".join(sorted(ai_additions))}')
        col_map.update(ai_additions)

    return col_map


def _get(row, col_map, field):
    """Safely retrieve a value from a row by field name."""
    idx = col_map.get(field)
    if idx is None or idx >= len(row):
        return None
    val = row[idx]
    if val is None:
        return None
    # Excel stores a typed-in phone number or ZIP as a number, and str() of
    # that float is "8005551234.0".
    if isinstance(val, float) and val.is_integer():
        val = int(val)
    val = str(val).strip()
    return val if val else None


def _split_full_name(full_name):
    """(first, last) from a single name cell. "Reed, Dana" is last-first;
    anything else is split at the final space, so a two-word first name
    stays together rather than a double-barrelled surname."""
    full_name = ' '.join(full_name.split())
    if ',' in full_name:
        last, _, first = full_name.partition(',')
        return first.strip() or None, last.strip() or None
    first, _, last = full_name.rpartition(' ')
    if not first:
        return full_name, None
    return first, last


def _build_lead_details(row, col_map):
    """Combine all notes fields into a single Lead Details string."""
    parts = []
    for field in ('notes_action', 'notes_body', 'survey_answers', 'qualifiers'):
        val = _get(row, col_map, field)
        if val:
            parts.append(val)
    return '\n\n'.join(parts) if parts else None


def row_to_lead(row, col_map):
    """
    One sheet row as a lead dict, by column map alone -- no skipping.

    parse_workbook() filters what this returns; spreadsheet_import.py uses it
    as-is as the fallback for rows its model pass could not read, where
    nothing but an explicit Cold may keep a row out.
    """
    first = _get(row, col_map, 'first_name')
    last = _get(row, col_map, 'last_name')
    full_name = _get(row, col_map, 'full_name')
    if full_name and not first and not last:
        first, last = _split_full_name(full_name)
    return {
        'first_name':   first,
        'last_name':    last,
        'organization': _get(row, col_map, 'organization'),
        'title':        _get(row, col_map, 'title'),
        'phone':        _get(row, col_map, 'phone') or _get(row, col_map, 'phone_mobile'),
        'email':        _get(row, col_map, 'email'),
        'address1':     _get(row, col_map, 'address1'),
        'address2':     _get(row, col_map, 'address2'),
        'city':         _get(row, col_map, 'city'),
        'state':        _get(row, col_map, 'state'),
        'zip':          _get(row, col_map, 'zip'),
        'country':      _get(row, col_map, 'country'),
        'rank':         _get(row, col_map, 'rank'),
        'lead_details': _build_lead_details(row, col_map),
    }


def parse_excel(file_path, report=None):
    """
    Read an Excel file and return a list of lead dicts.

    Each dict has these keys (value is None if not found):
      first_name, last_name, organization, title, phone, email,
      address1, address2, city, state, zip, country,
      lead_details, rank

    `report`, if given, is filled in as parse_workbook() describes.
    """
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f'Excel file not found: {file_path}')
    if path.suffix.lower() == '.xls':
        raise ValueError(
            'Legacy .xls workbooks are not supported. Save the workbook as .xlsx first.'
        )
    return parse_workbook(path, report=report)


def parse_workbook(source, report=None):
    """
    parse_excel() for a path *or* an open binary file.

    The dashboard's spreadsheet upload hands this the uploaded bytes wrapped
    in BytesIO, so a customer list is parsed without ever being written to
    disk -- the same rule manual_lead.py keeps for screenshots.

    `report`, if given, is a dict this fills in with what was skipped and
    why, so a caller can tell the operator rather than silently importing
    fewer rows than the sheet has:
      columns        -- the internal fields that were found
      header_row     -- 1-based sheet row the headers were read from
      skipped_cold   -- rows ranked Cold
      skipped_empty  -- rows with no name and no organization
    """
    wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
    try:
        # The sheet that was open when the workbook was saved -- usually the
        # one with the leads on it.
        ws = wb.active
        rows = _read_rows(ws) if ws else []
        header_idx, headers = _find_header_row(rows)
        col_map = _build_column_map(headers) if headers else {}
        if not _maps_name_or_company(col_map):
            # Saved while looking at a summary or pivot tab: try the others in
            # order, by alias only so a many-sheet workbook is not one model
            # call per sheet. Reading the open sheet regardless found nothing
            # and told the operator so, while the leads sat one tab over.
            for other in wb.worksheets:
                if other is ws:
                    continue
                other_rows = _read_rows(other)
                other_idx, other_headers = _find_header_row(other_rows)
                normalized = {str(h).strip().lower() for h in other_headers if h}
                if any(normalized & set(COLUMN_ALIASES[f]) for f in _NAME_OR_COMPANY_FIELDS):
                    rows, header_idx, headers = other_rows, other_idx, other_headers
                    col_map = _build_column_map(headers)
                    break
    finally:
        wb.close()

    # Report which fields were found
    found = sorted(col_map.keys())
    missing = [f for f in COLUMN_ALIASES if f not in col_map]
    print(f'  Columns mapped : {", ".join(found)}')
    if missing:
        print(f'  Columns missing: {", ".join(missing)} (those fields will be blank)')

    skipped_cold = skipped_empty = 0
    leads = []
    for idx in range(header_idx + 1, len(rows)):
        row = rows[idx]
        row_number = idx + 1

        # Skip completely empty rows
        if all(v is None or str(v).strip() == '' for v in row):
            continue

        lead = row_to_lead(row, col_map)

        # Skip rows with no name and no organization
        if not lead['first_name'] and not lead['last_name'] and not lead['organization']:
            skipped_empty += 1
            continue

        # Skip Cold leads — not entered into Oracle
        rank_val = lead['rank']
        if rank_val and rank_val.strip().lower() == 'cold':
            skipped_cold += 1
            continue

        lead['row_number'] = row_number
        leads.append(lead)

    if report is not None:
        report.update({
            'columns': found,
            'header_row': header_idx + 1,
            'skipped_cold': skipped_cold,
            'skipped_empty': skipped_empty,
        })
    return leads


def get_sheet_names(file_path):
    wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
    return wb.sheetnames


if __name__ == '__main__':
    import sys
    path = sys.argv[1] if len(sys.argv) > 1 else r'C:\Users\kkaess\Downloads\Leads for automation.xlsx'
    leads = parse_excel(path)
    print(f'\nParsed {len(leads)} leads. First 3:\n')
    for lead in leads[:3]:
        print(lead)
        print()
