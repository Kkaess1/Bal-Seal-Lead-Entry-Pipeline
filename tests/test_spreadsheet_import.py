"""
Importing a spreadsheet of trade-show leads uploaded on the dashboard.

Every lead list looks different, so the model reads it -- once for the
layout, then ten rows at a time -- and its answers are held to what the rows
actually say. What is pinned here:

  * every row becomes a lead unless it explicitly says Cold (Keaton,
    2026-09-28: err on the side of importing; a human filters on Pending
    Review);
  * each lead carries the show, so Oracle titles it
    "COMPANY - Contact - Tradeshow - Show";
  * nothing the model invents reaches a lead, and nothing on the sheet is lost;
  * an AI outage still imports every row.

The model is always faked: FakeModel answers the layout prompt and the row
prompt from fixtures, so no test makes a real OpenAI call.
"""

import codecs
import datetime
import io
import json
import os
import re
import tempfile
import threading
import unittest
import zipfile
from pathlib import Path
from unittest.mock import MagicMock, patch

import openpyxl

import db
import excel_parser
import lead_stages
import oracle_api
import pipeline_submit
import spreadsheet_import
from dashboard.app import app


def _save(workbook):
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _sheet(rows, title='Leads', workbook=None):
    """A workbook (or one more sheet on it) with `rows` written from A1.
    Each row is a list; None leaves a cell empty."""
    if workbook is None:
        workbook = openpyxl.Workbook()
        ws = workbook.active
        ws.title = title
    else:
        ws = workbook.create_sheet(title)
    for row in rows:
        ws.append(list(row))
    return workbook


# The shape of a real show report: a title block, a blank row, then headers.
SHOW_REPORT = _save(_sheet([
    ['Space Tech Expo 2026 - Lead Report'],
    [],
    ['Company', 'Contact', 'Email', 'Location', 'Rating', 'Booth notes'],
    ['Northwind Connectors', 'Sam Ortiz', 'sortiz@northwind-conn.com',
     'Irvine, CA 92618', 'Hot', 'Needs contacts for a satellite connector, 5k/yr'],
    ['Weber Antriebe GmbH', 'Jonas Weber', 'jonas@weber-antriebe.de',
     'Munich', 'Warm', 'Catalog request'],
]))

SHOW_LAYOUT = {
    'sheets': [{
        'sheet': 'Leads', 'has_leads': True, 'header_row': 3, 'first_data_row': 4,
        'columns': {'A': 'company', 'B': 'contact full name', 'C': 'email',
                    'D': 'city, state and ZIP combined', 'E': 'rating', 'F': 'notes'},
        'trade_show': 'Space Tech Expo',
    }],
    'trade_show': 'Space Tech Expo',
    'trade_show_evidence': 'row 1: "Space Tech Expo 2026 - Lead Report"',
}

SHOW_ROWS = {
    '1:4': {'id': '1:4', 'kind': 'lead', 'fields': {
        'first_name': 'Sam', 'last_name': 'Ortiz', 'organization': 'Northwind Connectors',
        'email': 'sortiz@northwind-conn.com', 'city': 'Irvine', 'state': 'CA', 'zip': '92618',
        'country': 'United States', 'rank': 'Hot', 'market': 'AEROSPACE/DEFENSE',
    }},
    '1:5': {'id': '1:5', 'kind': 'lead', 'fields': {
        'first_name': 'Jonas', 'last_name': 'Weber', 'organization': 'Weber Antriebe GmbH',
        'email': 'jonas@weber-antriebe.de', 'city': 'Munich', 'country': 'Germany',
        'rank': 'Warm', 'market': 'MOBILITY',
    }},
}


class FakeModel:
    """
    Stands in for ai_provider.call_llm. Answers the layout prompt with
    `layout` and each row prompt with the `rows` entries for the ids it was
    asked about. Either may be an exception to raise instead.
    """

    def __init__(self, layout=None, rows=None):
        self.layout = layout
        self.rows = rows or {}
        self.row_calls = []
        self.layout_calls = 0
        self.max_tokens = []
        self._lock = threading.Lock()

    def __call__(self, prompt, provider=None, max_tokens=None, model=None):
        with self._lock:
            self.max_tokens.append(max_tokens)
        if 'Work out its layout' in prompt:
            with self._lock:
                self.layout_calls += 1
            if isinstance(self.layout, Exception):
                raise self.layout
            return json.dumps(self.layout or {'sheets': []})
        block = prompt.split('<<<ROWS\n', 1)[1].split('\nROWS>>>', 1)[0]
        ids = [json.loads(line)['id'] for line in block.splitlines() if line.strip()]
        with self._lock:
            self.row_calls.append(ids)
        if isinstance(self.rows, Exception):
            raise self.rows
        return json.dumps({'rows': [self.rows[i] for i in ids if i in self.rows]})


class _DbTestCase(unittest.TestCase):
    def setUp(self):
        self.tempdir = tempfile.TemporaryDirectory()
        self.addCleanup(self.tempdir.cleanup)
        self.original_db_path = db.DB_PATH
        db.DB_PATH = Path(self.tempdir.name) / 'test.db'
        self.addCleanup(lambda: setattr(db, 'DB_PATH', self.original_db_path))
        db.init_db()
        # The fallback's column mapper has its own AI call and cache file;
        # the alias table covers every header these tests use.
        for target, value in (
            ('excel_parser.map_columns_with_ai', {}),
            ('spreadsheet_import.ai_provider.get_ai_provider', ('openai', 'test-key')),
        ):
            patcher = patch(target, return_value=value)
            patcher.start()
            self.addCleanup(patcher.stop)

    def _import(self, data=SHOW_REPORT, layout=SHOW_LAYOUT, rows=SHOW_ROWS,
                filename='leads.xlsx', **kwargs):
        model = FakeModel(layout, rows)
        with patch('spreadsheet_import.ai_provider.call_llm', side_effect=model):
            result = spreadsheet_import.import_spreadsheet(data, filename, **kwargs)
        return result, model

    def _leads(self, result):
        return [db.get_lead(i) for i in result['lead_ids']]


class ShowReportTests(_DbTestCase):
    def test_a_show_report_with_a_title_block_becomes_leads_titled_for_the_show(self):
        """The case the feature exists for. The Oracle project name is
        "COMPANY - Contact - Tradeshow - Show", with the show read off the
        sheet's own title row."""
        result, _ = self._import()
        sam, jonas = self._leads(result)

        self.assertEqual(sam['status'], 'pending_review')
        self.assertEqual(lead_stages.page_of(sam['status']), 'Pending Review')
        self.assertEqual(sam['source'], 'spreadsheet_upload')
        self.assertIsNone(sam['message_id'])
        self.assertEqual((sam['first_name'], sam['last_name']), ('Sam', 'Ortiz'))
        self.assertEqual((sam['city'], sam['state'], sam['zip']), ('Irvine', 'CA', '92618'))
        self.assertEqual(sam['rank'], 'Hot')
        self.assertEqual(sam['market'], 'AEROSPACE/DEFENSE')
        self.assertEqual(sam['trade_show'], 'Space Tech Expo')
        self.assertEqual(jonas['trade_show'], 'Space Tech Expo')

        source_label, show = pipeline_submit._project_label(sam)
        self.assertEqual(
            oracle_api.build_project_name(sam, show, source_label),
            'NORTHWIND CONNECTORS - Sam Ortiz - Tradeshow - Space Tech Expo',
        )
        self.assertEqual(oracle_api.submission_blockers(sam), [])

    def test_the_notes_record_where_each_lead_came_from(self):
        result, _ = self._import(filename='show.xlsx')
        notes = self._leads(result)[0]['extraction_notes']
        self.assertIn('show.xlsx', notes)
        self.assertIn('sheet "Leads", row 4', notes)
        self.assertEqual(self._leads(result)[0]['source_row_number'], 4)

    def test_every_unused_cell_is_kept_verbatim_in_lead_details(self):
        """Nothing on the sheet is lost. The booth note and the combined
        location cell go to Oracle word for word; the cells the fields used
        (company, name, email, rating) are not repeated."""
        result, _ = self._import()
        details = self._leads(result)[0]['lead_details']
        self.assertIn('Booth notes: Needs contacts for a satellite connector, 5k/yr', details)
        self.assertIn('Location: Irvine, CA 92618', details)
        self.assertNotIn('Northwind Connectors', details)
        self.assertNotIn('sortiz@northwind-conn.com', details)

    def test_rows_are_routed_to_a_territory_manager(self):
        result, _ = self._import()
        for lead in self._leads(result):
            self.assertTrue(lead['assigned_tm'], f'{lead["organization"]} was left unassigned')

    def test_rows_are_read_ten_at_a_time_and_given_room_to_reason(self):
        """Small batches are the accuracy choice. The pinned model's reasoning
        counts against max_output_tokens, and a batch cut off mid-JSON is ten
        leads that fall back to the column headers."""
        rows = [['Company', 'Contact']] + [[f'Co {i}', f'Person {i}'] for i in range(23)]
        layout = {'sheets': [{'sheet': 'Leads', 'has_leads': True,
                              'header_row': 1, 'first_data_row': 2}]}
        _, model = self._import(_save(_sheet(rows)), layout=layout, rows={})
        self.assertEqual(sorted(len(c) for c in model.row_calls), [3, 10, 10])
        self.assertTrue(all(t >= 16384 for t in model.max_tokens))


class ColdAndNotALeadTests(_DbTestCase):
    SHEET = _save(_sheet([
        ['Company', 'Contact', 'Rating', 'Notes'],
        ['Acme Seals', 'Ann Lee', 'Cold', 'Just browsing'],
        ['Brite Pumps', 'Bo Chan', 'Low', 'Maybe next year'],
        [None, None, None, 'Spoke with someone from a pump company, no card'],
        ['Crane Co', 'Cy Diaz', 'Warm', 'Needs cold-rolled steel springs'],
    ]))
    LAYOUT = {'sheets': [{'sheet': 'Leads', 'has_leads': True,
                          'header_row': 1, 'first_data_row': 2}],
              'trade_show': None}

    def test_a_row_is_skipped_only_when_it_explicitly_says_cold(self):
        """Keaton, 2026-09-28: unless a lead specifically says Cold it is a
        lead, and a human filters it on Pending Review. The model calling
        Brite Pumps cold ("Low") does not drop it, and a row with no name or
        company at all is still imported."""
        rows = {
            '1:2': {'id': '1:2', 'kind': 'cold', 'reason': 'Rating: Cold'},
            '1:3': {'id': '1:3', 'kind': 'cold', 'reason': 'Rating: Low',
                    'fields': {'organization': 'Brite Pumps'}},
            '1:4': {'id': '1:4', 'kind': 'lead', 'fields': {}},
            '1:5': {'id': '1:5', 'kind': 'lead', 'fields': {'organization': 'Crane Co'}},
        }
        result, _ = self._import(self.SHEET, layout=self.LAYOUT, rows=rows)
        self.assertEqual(result['skipped_cold'], ['Leads row 2 (Acme Seals)'])
        leads = self._leads(result)
        self.assertEqual([l['source_row_number'] for l in leads], [3, 4, 5])
        self.assertIn('does not say so', leads[0]['extraction_notes'])
        self.assertIn('no card', leads[1]['lead_details'])

    def test_a_note_about_cold_rolled_steel_is_not_a_cold_lead_without_the_ai(self):
        """With no model the fallback decides alone, and only a cell reading
        exactly "Cold" counts -- matching the word anywhere dropped Crane Co."""
        with patch('spreadsheet_import.ai_provider.get_ai_provider', return_value=(None, None)):
            result, model = self._import(self.SHEET, layout=self.LAYOUT, rows={})
        self.assertEqual(model.layout_calls, 0)
        self.assertEqual(len(result['skipped_cold']), 1)
        self.assertIn('Crane Co', [l['organization'] for l in self._leads(result)])

    def test_a_row_the_model_calls_not_a_lead_is_imported_unless_it_is_trivial(self):
        """A totals line is safe to skip. A row with an email address in it is
        not, whatever the model thinks of it -- it goes to Pending Review with
        the model's doubt attached."""
        data = _save(_sheet([
            ['Company', 'Contact', 'Email'],
            ['Acme Seals', 'Ann Lee', 'ann@acme.com'],
            ['Total leads:', 2],
            ['SEE BELOW', 'booth staff', 'boothstaff@balseal.com'],
        ]))
        rows = {
            '1:2': {'id': '1:2', 'kind': 'lead', 'fields': {'organization': 'Acme Seals'}},
            '1:3': {'id': '1:3', 'kind': 'not_a_lead', 'reason': 'Total leads: 2'},
            '1:4': {'id': '1:4', 'kind': 'not_a_lead', 'reason': 'staff note'},
        }
        result, _ = self._import(data, layout=self.LAYOUT, rows=rows)
        self.assertEqual(result['skipped_not_leads'], ['Leads row 3 (Total leads:)'])
        kept = self._leads(result)[-1]
        self.assertEqual(kept['source_row_number'], 4)
        self.assertIn('may not be a lead', kept['extraction_notes'])


class CheckingTheModelTests(_DbTestCase):
    SHEET = _save(_sheet([
        ['Company', 'Contact', 'Phone', 'City'],
        ['Reed Instruments', 'Dana Reed', '(949) 555-0142', 'Austin'],
    ]))
    LAYOUT = {'sheets': [{'sheet': 'Leads', 'has_leads': True,
                          'header_row': 1, 'first_data_row': 2}],
              'trade_show': None}

    def _one(self, fields, data=None, **entry):
        rows = {'1:2': dict({'id': '1:2', 'kind': 'lead', 'fields': fields}, **entry)}
        result, _ = self._import(data or self.SHEET, layout=self.LAYOUT, rows=rows)
        return self._leads(result)[0]

    def test_anything_the_model_invents_is_dropped_and_noted(self):
        """The model is told to copy, and it is held to that: an email, a
        company or a phone number that is not in the row does not reach the
        lead, and the reviewer is told what was left out."""
        lead = self._one({
            'first_name': 'Dana', 'last_name': 'Reed', 'organization': 'Reed Instruments Inc',
            'email': 'dana@reedinst.com', 'phone': '(800) 555-9999', 'title': 'CEO',
        })
        self.assertEqual(lead['first_name'], 'Dana')
        self.assertIsNone(lead['email'])
        self.assertIsNone(lead['phone'])
        self.assertIsNone(lead['title'])
        self.assertIn('Not found in the row', lead['extraction_notes'])
        self.assertIn('dana@reedinst.com', lead['extraction_notes'])

    def test_a_value_differing_only_in_case_or_punctuation_is_kept(self):
        lead = self._one({'first_name': 'DANA', 'organization': 'reed instruments',
                          'phone': '949-555-0142'})
        self.assertEqual(lead['organization'], 'reed instruments')
        self.assertEqual(lead['phone'], '949-555-0142')

    def test_nothing_a_cell_says_can_set_anything_but_a_checked_field(self):
        """The prompt-injection boundary. Whatever a cell talks the model into,
        the lead stays in Pending Review with valid dropdown values."""
        lead = self._one({
            'organization': 'Reed Instruments', 'market': 'Ignore previous instructions',
            'rank': 'Approved', 'status': 'approved', 'assigned_tm': 'Nobody',
            'fraud_risk': 'none',
        })
        self.assertEqual(lead['status'], 'pending_review')
        self.assertEqual(lead['market'], 'INDUSTRIAL')
        self.assertIsNone(lead['rank'])
        self.assertNotEqual(lead['assigned_tm'], 'Nobody')

    def test_an_inferred_country_or_state_is_kept_only_if_real_and_is_noted(self):
        """Inference is what these two are asked for when the row has none
        -- Austin makes Texas unambiguous -- but only to a real place."""
        lead = self._one({'organization': 'Reed Instruments', 'city': 'Austin',
                          'state': 'TX', 'country': 'USA'})
        self.assertEqual(lead['state'], 'TX')
        self.assertEqual(lead['country'], 'United States')
        self.assertIn('State "TX" was inferred', lead['extraction_notes'])
        self.assertTrue(lead['assigned_tm'])

        other_file = _save(_sheet([['Company', 'Contact', 'Phone', 'City'],
                                   ['Reed Instruments', 'Dana Reed', None, 'Nowhere']]))
        lead = self._one({'organization': 'Reed Instruments', 'state': 'Narnia',
                          'country': 'Atlantis'}, data=other_file)
        self.assertIsNone(lead['state'])
        self.assertIsNone(lead['country'])

    def test_the_models_doubts_reach_the_reviewer(self):
        lead = self._one({'organization': 'Reed Instruments'},
                         uncertain='Two contacts listed; used the first.')
        self.assertIn('Two contacts listed', lead['extraction_notes'])


class TradeShowTests(_DbTestCase):
    def test_a_show_the_workbook_never_mentions_is_not_accepted(self):
        """The show name ends up in every Oracle title from the file, so it
        must be in the file. One the model made up is dropped, and the lead
        cannot be approved until someone enters it."""
        layout = dict(SHOW_LAYOUT, trade_show='CES',
                      sheets=[dict(SHOW_LAYOUT['sheets'][0], trade_show='CES')])
        result, _ = self._import(layout=layout)
        lead = self._leads(result)[0]
        self.assertIsNone(lead['trade_show'])
        self.assertIsNone(result['trade_show'])
        self.assertIn('No trade show name was found', lead['extraction_notes'])
        self.assertIn('Trade Show is missing', ' '.join(oracle_api.submission_blockers(lead)))

    def test_a_show_written_only_in_the_filename_is_found(self):
        data = _save(_sheet([['Company', 'Contact'], ['Acme Seals', 'Ann Lee']]))
        layout = {'sheets': [{'sheet': 'Leads', 'has_leads': True, 'header_row': 1,
                              'first_data_row': 2}],
                  'trade_show': 'Space Tech Expo'}
        result, _ = self._import(data, layout=layout, rows={
            '1:2': {'id': '1:2', 'kind': 'lead', 'fields': {'organization': 'Acme Seals'}}},
            filename='SpaceTechExpo2026.xlsx')
        self.assertEqual(self._leads(result)[0]['trade_show'], 'Space Tech Expo')

    def test_a_row_naming_its_own_show_keeps_it(self):
        """A season's list with an Event column carries a show per row."""
        data = _save(_sheet([['Company', 'Event'], ['Acme Seals', 'MD&M West'],
                             ['Brite Pumps', 'Space Tech Expo']]))
        layout = {'sheets': [{'sheet': 'Leads', 'has_leads': True, 'header_row': 1,
                              'first_data_row': 2}], 'trade_show': None}
        rows = {f'1:{n}': {'id': f'1:{n}', 'kind': 'lead',
                           'fields': {'organization': org, 'trade_show': show}}
                for n, org, show in ((2, 'Acme Seals', 'MD&M West'),
                                     (3, 'Brite Pumps', 'Space Tech Expo'))}
        result, _ = self._import(data, layout=layout, rows=rows)
        self.assertEqual([l['trade_show'] for l in self._leads(result)],
                         ['MD&M West', 'Space Tech Expo'])

    def test_the_forms_show_is_used_only_when_the_workbook_names_none(self):
        result, _ = self._import(fallback_trade_show='Wrong Show')
        self.assertEqual(self._leads(result)[0]['trade_show'], 'Space Tech Expo')

        data = _save(_sheet([['Company'], ['Acme Seals']]))
        layout = {'sheets': [{'sheet': 'Leads', 'has_leads': True, 'header_row': 1,
                              'first_data_row': 2}], 'trade_show': None}
        result, _ = self._import(data, layout=layout, rows={},
                                 fallback_trade_show='  Sensors   Expo ')
        self.assertEqual(self._leads(result)[0]['trade_show'], 'Sensors Expo')
        self.assertTrue(result['trade_show_from_form'])

    def test_a_trade_show_typed_on_any_lead_titles_it_for_the_show(self):
        """The field is on the review form for every lead; filling it in is
        how a reviewer says "this one was met at a show"."""
        self.assertEqual(
            pipeline_submit._project_label({'source': 'email_ai_extract',
                                            'trade_show': 'MD&M West'}),
            ('Tradeshow', 'MD&M West'),
        )
        self.assertEqual(
            pipeline_submit._project_label({'source': 'email_ai_extract'}),
            ('Inquiry', 'Website/Email'),
        )

    def test_the_trade_show_field_is_on_the_review_form(self):
        self.assertIn('trade_show', db.LEAD_SCHEMA_FIELDS)
        from dashboard.app import LEAD_FIELD_GROUPS
        self.assertIn('trade_show', [f for _, fields in LEAD_FIELD_GROUPS for f in fields])


class LeadSourceTests(_DbTestCase):
    """Keaton, 2026-09-28: a lead from a trade show goes to Oracle with
    Lead Source "Trade Shows", not "Website"; and an upload asks where its
    leads are from."""

    def test_a_trade_show_upload_is_sent_as_trade_shows(self):
        result, _ = self._import()
        lead = self._leads(result)[0]
        self.assertEqual(lead['lead_source'], 'Trade Shows')
        self.assertEqual(oracle_api.lead_source_for(lead), 'Trade Shows')

    def test_any_lead_from_a_show_is_trade_shows_and_everything_else_website(self):
        self.assertEqual(oracle_api.lead_source_for({'trade_show': 'MD&M West'}), 'Trade Shows')
        # A sheet attached to an email has always been titled a trade show's.
        self.assertEqual(oracle_api.lead_source_for({'source': 'email_excel_attachment'}),
                         'Trade Shows')
        self.assertEqual(oracle_api.lead_source_for({'source': 'email_ai_extract'}), 'Website')
        self.assertEqual(oracle_api.lead_source_for({'source': 'manual_entry'}), 'Website')

    def test_a_source_chosen_for_the_lead_wins_and_a_made_up_one_does_not(self):
        self.assertEqual(oracle_api.lead_source_for(
            {'trade_show': 'MD&M West', 'lead_source': 'Referrals'}), 'Referrals')
        self.assertEqual(oracle_api.lead_source_for(
            {'trade_show': 'MD&M West', 'lead_source': 'Carrier pigeon'}), 'Trade Shows')

    def test_the_submitter_sends_each_leads_own_source(self):
        """The client's configured source is only the default -- it used to
        be sent for every lead, which is how trade-show leads went out as
        Website."""
        client = oracle_api.OracleLeadApiClient(
            {'base_url': 'https://fake.oracle.example', 'lead_source': 'Website',
             'api_timeout_seconds': 10}, username='u', password='p', session=MagicMock())
        seen = {}

        def capture(lead, trade_show_name, **kwargs):
            seen['lead_source'] = kwargs['lead_source']
            raise RuntimeError('stop here')

        with patch.object(client, 'resolve_existing_party', return_value=None), \
             patch('oracle_api.build_lead_payload', side_effect=capture):
            with self.assertRaises(RuntimeError):
                client.submit_lead({'trade_show': 'Space Tech Expo'}, 'Space Tech Expo')
        self.assertEqual(seen['lead_source'], 'Trade Shows')

    def test_a_list_from_elsewhere_carries_its_source_and_no_show(self):
        """A referral list is not a trade show's: no show name, no show
        required, and its Oracle title names the source instead."""
        result, _ = self._import(lead_source='Referrals')
        lead = self._leads(result)[0]
        self.assertEqual(lead['lead_source'], 'Referrals')
        self.assertIsNone(lead['trade_show'])
        self.assertNotIn('No trade show name', lead['extraction_notes'])
        self.assertEqual(oracle_api.submission_blockers(lead), [])
        source_label, show = pipeline_submit._project_label(lead)
        self.assertEqual(oracle_api.build_project_name(lead, show, source_label),
                         'NORTHWIND CONNECTORS - Sam Ortiz - Referrals')
        message, _ = spreadsheet_import.summarize(result)
        self.assertIn('lead source: Referrals', message)
        self.assertNotIn('trade show', message.lower())

    def test_an_unknown_lead_source_is_refused(self):
        with self.assertRaisesRegex(spreadsheet_import.ImportError_, 'where these leads are from'):
            self._import(lead_source='Carrier pigeon')
        with self.assertRaises(spreadsheet_import.ImportError_):
            spreadsheet_import.start_import(SHOW_REPORT, 'leads.xlsx', lead_source='')
        self.assertEqual(db.count_leads(), 0)


class CsvTests(_DbTestCase):
    """Badge-scanner and event-app exports are often CSV only. A CSV is read
    as a workbook of one sheet named after the file, and from there is
    handled exactly like an .xlsx."""

    LAYOUT = {'sheets': [{'sheet': 'scans', 'has_leads': True,
                          'header_row': 2, 'first_data_row': 3}],
              'trade_show': 'Space Tech Expo'}

    def _rows(self, data, filename='scans.csv'):
        return spreadsheet_import.read_csv(data, filename)['sheets'][0]['rows']

    def test_a_csv_show_export_becomes_leads_like_a_workbook(self):
        data = (
            'Space Tech Expo 2026 lead export\r\n'
            'Company,Contact,ZIP,Notes\r\n'
            'Northwind Connectors,Sam Ortiz,02134,"Wants a quote, 5k/yr"\r\n'
        ).encode('utf-8')
        rows = {'1:3': {'id': '1:3', 'kind': 'lead', 'fields': {
            'organization': 'Northwind Connectors', 'first_name': 'Sam',
            'last_name': 'Ortiz', 'zip': '02134'}}}
        result, model = self._import(data, layout=self.LAYOUT, rows=rows, filename='scans.csv')
        self.assertEqual(model.row_calls, [['1:3']])
        lead = self._leads(result)[0]
        self.assertEqual(lead['organization'], 'Northwind Connectors')
        # Already text in a CSV, so the leading zero survives.
        self.assertEqual(lead['zip'], '02134')
        self.assertEqual(lead['trade_show'], 'Space Tech Expo')
        self.assertIn('Notes: Wants a quote, 5k/yr', lead['lead_details'])
        self.assertIn('sheet "scans", row 3', lead['extraction_notes'])

    def test_a_european_semicolon_export_is_split_on_semicolons(self):
        """Excel writes ";" wherever the decimal separator is a comma. Commas
        inside the notes outnumber the semicolons, but only the semicolon
        splits every row the same way."""
        data = ('Firma;Name;Notiz\n'
                'Weber Antriebe GmbH;Jonas Weber;Katalog, Muster, Preise\n'
                'Helio Dynamics;Anika Rao;Drehdichtung, 40 mm, 5000 U/min\n').encode('utf-8')
        self.assertEqual(self._rows(data)[1],
                         ['Weber Antriebe GmbH', 'Jonas Weber', 'Katalog, Muster, Preise'])

    def test_excels_sep_hint_line_is_obeyed_and_not_read_as_a_row(self):
        data = 'sep=|\nCompany|Contact\nAcme, Inc.|Ann Lee\n'.encode('utf-8')
        self.assertEqual(self._rows(data), [['Company', 'Contact'], ['Acme, Inc.', 'Ann Lee']])

    def test_every_encoding_excel_writes_is_read_with_its_accents(self):
        """"CSV UTF-8" has a byte-order mark, plain "CSV" is the Windows code
        page, "Unicode Text" is UTF-16 with tabs. Reading the code page as
        UTF-8 failed on the first accented name; reading a BOM as text put it
        in front of the first header."""
        text = 'Company,Contact\nMüller Dichtungen,Zoë Lefèvre\n'
        for data in (
            codecs.BOM_UTF8 + text.encode('utf-8'),
            text.encode('cp1252'),
            text.replace(',', '\t').encode('utf-16'),
        ):
            rows = self._rows(data)
            self.assertEqual(rows[0][0], 'Company')
            self.assertEqual(rows[1], ['Müller Dichtungen', 'Zoë Lefèvre'])

    def test_a_binary_file_named_csv_is_refused_with_a_sentence(self):
        with self.assertRaisesRegex(spreadsheet_import.ImportError_, 'could not be read'):
            spreadsheet_import.import_spreadsheet(SHOW_REPORT, 'leads.csv')
        self.assertEqual(db.count_leads(), 0)

    def test_an_empty_csv_is_refused(self):
        with self.assertRaisesRegex(spreadsheet_import.ImportError_, 'no data'):
            spreadsheet_import.import_spreadsheet(b'\r\n,,\r\n', 'leads.csv')


class LayoutTests(_DbTestCase):
    def test_a_summary_tab_is_skipped_but_a_tab_with_emails_never_is(self):
        """Err toward importing: the model may call a tab "not leads", but one
        holding email addresses is read regardless."""
        workbook = _sheet([['Company', 'Contact'], ['Acme Seals', 'Ann Lee']], title='Day 1')
        _sheet([['Leads collected', 2]], title='Summary', workbook=workbook)
        _sheet([['Company', 'Email'], ['Brite Pumps', 'bo@brite.com']],
               title='Day 2', workbook=workbook)
        layout = {'sheets': [
            {'sheet': 'Day 1', 'has_leads': True, 'header_row': 1, 'first_data_row': 2},
            {'sheet': 'Summary', 'has_leads': False},
            {'sheet': 'Day 2', 'has_leads': False},
        ], 'trade_show': None}
        rows = {
            '1:2': {'id': '1:2', 'kind': 'lead', 'fields': {'organization': 'Acme Seals'}},
            '3:2': {'id': '3:2', 'kind': 'lead', 'fields': {'organization': 'Brite Pumps'}},
        }
        result, _ = self._import(_save(workbook), layout=layout, rows=rows)
        self.assertEqual(result['sheets'], ['Day 1', 'Day 2'])
        self.assertEqual([l['organization'] for l in self._leads(result)],
                         ['Acme Seals', 'Brite Pumps'])

    def test_a_repeated_header_and_a_duplicated_tab_are_not_imported_twice(self):
        workbook = _sheet([['Company', 'Contact'], ['Acme Seals', 'Ann Lee'],
                           ['Company', 'Contact'], ['Brite Pumps', 'Bo Chan']], title='All')
        _sheet([['Company', 'Contact'], ['Acme Seals', 'Ann Lee']], title='Day 1',
               workbook=workbook)
        layout = {'sheets': [
            {'sheet': 'All', 'has_leads': True, 'header_row': 1, 'first_data_row': 2},
            {'sheet': 'Day 1', 'has_leads': True, 'header_row': 1, 'first_data_row': 2},
        ]}
        result, model = self._import(_save(workbook), layout=layout, rows={})
        self.assertEqual(len(result['lead_ids']), 2)
        self.assertEqual(result['repeated_rows'], 1)
        self.assertEqual(sorted(i for c in model.row_calls for i in c), ['1:2', '1:4'])

    def test_data_away_from_a1_keeps_its_real_row_numbers(self):
        """A read-only sheet starts iterating at its declared dimension, so a
        list starting at B3 came back numbered from 1 and every "row N" in
        the notes pointed at the wrong line."""
        workbook = openpyxl.Workbook()
        ws = workbook.active
        ws.title = 'Leads'
        ws['B3'], ws['C3'] = 'Company', 'Contact'
        ws['B4'], ws['C4'] = 'Acme Seals', 'Ann Lee'
        layout = {'sheets': [{'sheet': 'Leads', 'has_leads': True, 'header_row': 3,
                              'first_data_row': 4}]}
        result, model = self._import(_save(workbook), layout=layout, rows={})
        self.assertEqual(model.row_calls, [['1:4']])
        self.assertEqual(self._leads(result)[0]['source_row_number'], 4)

    def test_a_zip_formatted_with_leading_zeros_keeps_them(self):
        """Excel's own Zip Code format stores 02134 as the number 2134."""
        workbook = openpyxl.Workbook()
        ws = workbook.active
        ws.append(['Company', 'ZIP', 'Scanned'])
        ws.append(['Acme Seals', 2134, datetime.datetime(2026, 3, 1)])
        ws['B2'].number_format = '00000'
        sheet = spreadsheet_import.read_workbook(_save(workbook))['sheets'][0]
        self.assertEqual(sheet['rows'][1], ['Acme Seals', '02134', '2026-03-01'])

    def test_a_layout_outage_falls_back_to_the_header_rules(self):
        result, model = self._import(
            _save(_sheet([['Company', 'Contact', 'Email'], ['Acme Seals', 'Ann Lee', 'a@acme.com']])),
            layout=RuntimeError('timeout'),
            rows={'1:2': {'id': '1:2', 'kind': 'lead', 'fields': {'organization': 'Acme Seals'}}},
        )
        self.assertEqual(model.row_calls, [['1:2']])
        self.assertEqual(self._leads(result)[0]['organization'], 'Acme Seals')


class OutageTests(_DbTestCase):
    SHEET = _save(_sheet([
        ['First Name', 'Last Name', 'Company', 'Email', 'State', 'Rank'],
        ['Dana', 'Reed', 'Reed Instruments', 'dana@reedinst.com', 'TX', 'Hot'],
        ['Ann', 'Lee', 'Acme Seals', 'ann@acme.com', 'TX', 'Cold'],
    ]))

    def test_an_ai_outage_still_imports_every_row_from_the_column_headers(self):
        """An operator who has uploaded a show's leads must never lose them to
        an API outage. The rows come in by the header rules, each saying so."""
        result, _ = self._import(self.SHEET, layout=RuntimeError('down'),
                                 rows=RuntimeError('down'))
        self.assertEqual(len(result['lead_ids']), 1)
        self.assertEqual(len(result['skipped_cold']), 1)
        lead = self._leads(result)[0]
        self.assertEqual((lead['first_name'], lead['organization']), ('Dana', 'Reed Instruments'))
        self.assertEqual(lead['rank'], 'Hot')
        self.assertEqual(lead['market'], 'INDUSTRIAL')
        self.assertIn('column headers only', lead['extraction_notes'])
        self.assertEqual(result['fallback_rows'], ['Leads row 2 (Dana Reed)'])

    def test_without_the_ai_a_summary_tab_is_not_imported_as_leads(self):
        """With no layout, nobody says which tab is a summary. One with no
        email and no name or company header imported its totals as leads."""
        workbook = _sheet([['First Name', 'Company'], ['Dana', 'Reed Co']], title='Leads')
        _sheet([['A leads', 1], ['B leads', 1]], title='Summary', workbook=workbook)
        result, _ = self._import(_save(workbook), layout=RuntimeError('down'),
                                 rows=RuntimeError('down'))
        self.assertEqual(result['sheets'], ['Leads'])
        self.assertEqual(len(result['lead_ids']), 1)

    def test_a_row_the_model_left_out_falls_back_rather_than_vanishing(self):
        layout = {'sheets': [{'sheet': 'Leads', 'has_leads': True, 'header_row': 1,
                              'first_data_row': 2}]}
        result, _ = self._import(self.SHEET, layout=layout, rows={})
        self.assertEqual(len(result['lead_ids']), 1)
        self.assertEqual(len(result['fallback_rows']), 1)


class SafetyTests(_DbTestCase):
    def test_uploading_the_same_file_twice_creates_no_duplicates(self):
        """A double-click, a browser retry or an operator unsure whether it
        worked must not put every lead on the page twice."""
        self._import()
        second, model = self._import()
        self.assertEqual(second['lead_ids'], [])
        self.assertEqual(second['already_imported'], 2)
        self.assertEqual(db.count_leads(), 2)
        self.assertEqual(model.row_calls, [])

    def test_a_second_upload_that_lands_during_the_analysis_is_reported_not_claimed(self):
        """The duplicate check runs before minutes of model calls. A second
        upload of the same file finishing in that window must be counted as
        already imported -- not handed back as new, and not a crash on the
        unique ingest_key index."""
        model = FakeModel(SHOW_LAYOUT, SHOW_ROWS)
        original = spreadsheet_import.read_rows
        raced = []

        def slow_read(rows, show, progress=None):
            found = original(rows, show, progress)
            if not raced:       # the second upload, once, mid-analysis
                raced.append(True)
                spreadsheet_import.import_spreadsheet(SHOW_REPORT, 'leads.xlsx')
            return found

        with patch('spreadsheet_import.ai_provider.call_llm', side_effect=model), \
             patch('spreadsheet_import.read_rows', side_effect=slow_read):
            result = spreadsheet_import.import_spreadsheet(SHOW_REPORT, 'leads.xlsx')
        self.assertEqual(result['lead_ids'], [])
        self.assertEqual(result['already_imported'], 2)
        self.assertEqual(db.count_leads(), 2)

    def test_an_import_that_fails_partway_leaves_nothing_behind(self):
        """One transaction per workbook: half a list is worse than none."""
        real_insert = db.insert_lead
        calls = []

        def fail_second(*args, **kwargs):
            calls.append(1)
            if len(calls) == 2:
                raise RuntimeError('disk full')
            return real_insert(*args, **kwargs)

        with patch('spreadsheet_import.db.insert_lead', side_effect=fail_second):
            with self.assertRaises(RuntimeError):
                self._import()
        self.assertEqual(db.count_leads(), 0)

    def test_a_file_that_is_not_a_workbook_is_refused_with_a_sentence(self):
        with self.assertRaisesRegex(spreadsheet_import.ImportError_, 'could not be read'):
            spreadsheet_import.import_spreadsheet(b'not a zip', 'leads.xlsx')

    def test_a_legacy_xls_is_refused_with_how_to_fix_it(self):
        for name in ('leads.xls', 'leads.numbers'):
            with self.assertRaisesRegex(spreadsheet_import.ImportError_, 'save it as'):
                spreadsheet_import.import_spreadsheet(SHOW_REPORT, name)

    def test_an_empty_workbook_is_refused(self):
        with self.assertRaisesRegex(spreadsheet_import.ImportError_, 'no data'):
            spreadsheet_import.import_spreadsheet(_save(openpyxl.Workbook()), 'e.xlsx')

    def test_a_workbook_over_the_row_limit_is_refused_before_any_row_is_read(self):
        with patch.object(spreadsheet_import, 'MAX_ROWS', 1):
            with self.assertRaisesRegex(spreadsheet_import.ImportError_, 'at most 1'):
                _, model = self._import()
        self.assertEqual(db.count_leads(), 0)

    def test_a_spreadsheet_lead_without_a_show_cannot_be_approved(self):
        lead = {'source': 'spreadsheet_upload', 'lead_source': 'Trade Shows',
                'country': 'US', 'state': 'TX', 'assigned_tm': 'Someone'}
        self.assertIn('Trade Show is missing', ' '.join(oracle_api.submission_blockers(lead)))
        lead['trade_show'] = 'Space Tech Expo'
        self.assertEqual(oracle_api.submission_blockers(lead), [])
        # Every other source is unaffected.
        self.assertEqual(oracle_api.submission_blockers(dict(lead, source='manual_entry',
                                                             trade_show=None)), [])


class BackgroundJobTests(_DbTestCase):
    def test_a_job_reports_progress_and_its_result(self):
        model = FakeModel(SHOW_LAYOUT, SHOW_ROWS)
        with patch('spreadsheet_import.ai_provider.call_llm', side_effect=model):
            job_id = spreadsheet_import.start_import(SHOW_REPORT, 'leads.xlsx')
            job = spreadsheet_import.wait_for_job(job_id, timeout=30)
        self.assertEqual(job['state'], 'done')
        self.assertEqual(len(job['result']['lead_ids']), 2)
        self.assertNotIn('_thread', job)
        message, ok = spreadsheet_import.summarize(job['result'])
        self.assertTrue(ok)
        self.assertIn('Imported 2 leads', message)
        self.assertIn('Space Tech Expo', message)

    def test_a_wrong_file_is_refused_before_a_job_starts(self):
        with self.assertRaises(spreadsheet_import.ImportError_):
            spreadsheet_import.start_import(b'nope', 'leads.xlsx')


class ExcelParserTests(unittest.TestCase):
    """The emailed-attachment path shares excel_parser; these hold it."""

    def test_a_workbook_can_be_parsed_from_memory(self):
        data = _save(_sheet([['First Name', 'Last Name', 'Company'], ['Dana', 'Reed', 'Reed Co']]))
        with patch('excel_parser.map_columns_with_ai', return_value={}):
            rows = excel_parser.parse_workbook(io.BytesIO(data))
        self.assertEqual(rows[0]['organization'], 'Reed Co')
        self.assertEqual(rows[0]['row_number'], 2)

    def test_a_phone_or_zip_typed_as_a_number_does_not_gain_a_decimal(self):
        data = _save(_sheet([['First Name', 'Last Name', 'Company', 'Phone', 'Zip'],
                             ['Dana', 'Reed', 'Co', 8005551234.0, 92610.0]]))
        with patch('excel_parser.map_columns_with_ai', return_value={}):
            row = excel_parser.parse_workbook(io.BytesIO(data))[0]
        self.assertEqual(row['phone'], '8005551234')
        self.assertEqual(row['zip'], '92610')

    def test_cold_rows_in_an_unnamed_rank_column_are_skipped_without_a_dimension_record(self):
        """Some exporters omit <dimension>, and openpyxl then returns ragged
        rows: a header row ending at column 8 had no index 8, the unnamed
        rank column was never recognised, and Cold rows came in as leads."""
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.append(['First Name', 'Last Name', 'Company', 'Title', 'Email',
                      'Phone', 'City', 'State'])
        sheet.append(['Dana', 'Reed', 'Reed Co', 'Eng', 'd@r.com', None, 'Austin', 'TX', 'Hot'])
        sheet.append(['Cole', 'Day', 'Nope Inc', 'Eng', 'c@n.com', None, 'Austin', 'TX', 'Cold'])
        source = zipfile.ZipFile(io.BytesIO(_save(workbook)))
        stripped = io.BytesIO()
        with zipfile.ZipFile(stripped, 'w') as target:
            for item in source.infolist():
                content = source.read(item.filename)
                if item.filename.startswith('xl/worksheets/'):
                    content = re.sub(rb'<dimension[^>]*/>', b'', content)
                target.writestr(item, content)
        with patch('excel_parser.map_columns_with_ai', return_value={}):
            report = {}
            rows = excel_parser.parse_workbook(io.BytesIO(stripped.getvalue()), report=report)
        self.assertEqual([r['organization'] for r in rows], ['Reed Co'])
        self.assertEqual(report['skipped_cold'], 1)

    def test_the_ai_mapper_cannot_reuse_a_column_the_aliases_claimed(self):
        """Asked about every unmapped field, the model can offer the Company
        column as full_name as well, and every company-only row would gain a
        contact named after its employer."""
        with patch('excel_parser.map_columns_with_ai',
                   return_value={'full_name': 0, 'phone': 2}):
            mapping = excel_parser._build_column_map(['Company', 'Email', 'Tel No.'])
        self.assertEqual(mapping['organization'], 0)
        self.assertNotIn('full_name', mapping)
        self.assertEqual(mapping['phone'], 2)

    def test_a_single_name_column_is_split_into_first_and_last(self):
        data = _save(_sheet([['Name', 'Company'], ['Dana Reed', 'Reed Co'],
                             ['Weber, Jonas', 'Weber AG']]))
        with patch('excel_parser.map_columns_with_ai', return_value={}):
            rows = excel_parser.parse_workbook(io.BytesIO(data))
        self.assertEqual([(r['first_name'], r['last_name']) for r in rows],
                         [('Dana', 'Reed'), ('Jonas', 'Weber')])


class SpreadsheetImportRouteTests(_DbTestCase):
    def setUp(self):
        super().setUp()
        app.config['TESTING'] = True
        self.client = app.test_client()
        with self.client.session_transaction() as sess:
            sess['_csrf_token'] = 'test-token'

    def _post(self, data, filename='leads.xlsx', **form):
        # The form's dropdown always sends a source; Trade Shows is its default.
        return self.client.post('/leads/import', data=dict({
            '_csrf_token': 'test-token',
            'workbook': (io.BytesIO(data), filename),
            'lead_source': 'Trade Shows',
        }, **form), content_type='multipart/form-data')

    def test_pending_review_offers_a_way_to_import_one(self):
        with patch('dashboard.app.db.get_lead_email_thread', return_value=[]):
            page = self.client.get('/').get_data(as_text=True)
        self.assertIn('/leads/import', page)
        self.assertIn('Import spreadsheet', page)

    def test_the_form_renders_with_the_show_fallback(self):
        page = self.client.get('/leads/import').get_data(as_text=True)
        self.assertIn('name="workbook"', page)
        self.assertIn('name="trade_show"', page)

    def test_the_form_asks_where_the_leads_are_from_defaulting_to_trade_shows(self):
        """Every Oracle Lead Source is offered, and a trade-show list -- what
        this upload is almost always for -- needs no change."""
        page = self.client.get('/leads/import').get_data(as_text=True)
        self.assertIn('name="lead_source"', page)
        for value in oracle_api.LEAD_SOURCE_VALUES:
            self.assertIn(f'<option value="{value}"', page)
        self.assertIn('<option value="Trade Shows" selected>', page)

    def test_a_lead_source_oracle_does_not_know_is_refused(self):
        response = self._post(SHOW_REPORT, lead_source='Carrier pigeon')
        self.assertEqual(response.status_code, 302)
        self.assertIn('/leads/import', response.headers['Location'])
        self.assertIn('Choose where these leads are from',
                      self.client.get('/leads/import').get_data(as_text=True))
        self.assertEqual(db.count_leads(), 0)

    def test_the_review_form_offers_the_lead_source_and_refuses_a_made_up_one(self):
        result, _ = self._import()
        lead_id = result['lead_ids'][0]
        with patch('dashboard.app.db.get_lead_email_thread', return_value=[]):
            detail = self.client.get(f'/leads/{lead_id}').get_data(as_text=True)
        self.assertIn('<select name="lead_source">', detail)
        self.assertRegex(detail, r'<option value="Trade Shows"\s+selected>')

        refused = self.client.post(f'/leads/{lead_id}/save', data={
            '_csrf_token': 'test-token', 'lead_source': 'Carrier pigeon'})
        self.assertEqual(refused.status_code, 400)
        self.client.post(f'/leads/{lead_id}/save', data={
            '_csrf_token': 'test-token', 'lead_source': 'Referrals'})
        self.assertEqual(db.get_lead(lead_id)['lead_source'], 'Referrals')

    def test_an_upload_runs_in_the_background_and_lands_on_the_first_new_card(self):
        """Minutes of model calls outlast a browser request, so the upload
        starts a job and the progress page polls it to the finish."""
        model = FakeModel(SHOW_LAYOUT, SHOW_ROWS)
        with patch('spreadsheet_import.ai_provider.call_llm', side_effect=model):
            response = self._post(SHOW_REPORT)
            self.assertEqual(response.status_code, 302)
            progress_url = response.headers['Location']
            job_id = progress_url.rstrip('/').rsplit('/', 1)[-1]
            spreadsheet_import.wait_for_job(job_id, timeout=30)

        page = self.client.get(progress_url).get_data(as_text=True)
        self.assertIn('Importing leads.xlsx', page)
        status = self.client.get(f'/leads/import/{job_id}/status').get_json()
        self.assertEqual(status['state'], 'done')

        finish = self.client.get(f'/leads/import/{job_id}/finish')
        self.assertEqual(finish.status_code, 302)
        self.assertIn('lead-card-1', finish.headers['Location'])
        self.assertEqual(db.count_leads(), 2)

    def test_the_forms_show_name_reaches_the_import(self):
        seen = {}

        def capture(data, filename, reviewed_by=None, fallback_trade_show=None,
                    lead_source=None):
            seen.update(show=fallback_trade_show, source=lead_source)
            return 'job'

        with patch('dashboard.app.spreadsheet_import.start_import', side_effect=capture):
            self._post(SHOW_REPORT, trade_show='Sensors Expo', lead_source='Trade Shows')
        self.assertEqual(seen, {'show': 'Sensors Expo', 'source': 'Trade Shows'})

    def test_a_bad_file_goes_back_to_the_form_with_the_reason(self):
        response = self._post(b'not a workbook')
        self.assertEqual(response.status_code, 302)
        self.assertIn('/leads/import', response.headers['Location'])
        page = self.client.get('/leads/import').get_data(as_text=True)
        self.assertIn('could not be read', page)
        self.assertEqual(db.count_leads(), 0)

    def test_an_unknown_job_goes_back_to_the_form(self):
        response = self.client.get('/leads/import/nosuchjob')
        self.assertEqual(response.status_code, 302)
        self.assertEqual(self.client.get('/leads/import/nosuchjob/status').status_code, 404)

    def test_a_csv_upload_is_accepted(self):
        data = 'Company,Contact\nAcme Seals,Ann Lee\n'.encode('utf-8')
        def start(data, filename, **kwargs):
            spreadsheet_import.check_file(data, filename)    # raises if refused
            return 'job'

        with patch('dashboard.app.spreadsheet_import.start_import', side_effect=start):
            response = self._post(data, filename='scans.csv')
        self.assertEqual(response.status_code, 302)
        self.assertIn('/leads/import/job', response.headers['Location'])

    def test_the_form_accepts_csv_files(self):
        page = self.client.get('/leads/import').get_data(as_text=True)
        self.assertIn('.csv', page)

    def test_a_large_upload_is_never_spooled_to_a_temp_file(self):
        """Werkzeug writes any upload over 500 KB to %TEMP% while parsing the
        form. For a customer list that is exactly the copy on disk the
        feature promises never to make."""
        with patch('tempfile.SpooledTemporaryFile.rollover') as rollover:
            response = self._post(os.urandom(600_000))
        self.assertEqual(response.status_code, 302)
        rollover.assert_not_called()

    def test_a_post_without_the_csrf_token_is_refused(self):
        response = self.client.post('/leads/import', data={
            'workbook': (io.BytesIO(SHOW_REPORT), 'leads.xlsx'),
        }, content_type='multipart/form-data')
        self.assertEqual(response.status_code, 400)
        self.assertEqual(db.count_leads(), 0)

    def test_an_imported_lead_renders_with_its_trade_show_editable(self):
        """A lead with no Outlook message behind it has to survive every page
        that assumes one, and the show must be correctable on the form."""
        result, _ = self._import()
        with patch('dashboard.app.db.get_lead_email_thread', return_value=[]):
            detail = self.client.get(f"/leads/{result['lead_ids'][0]}").get_data(as_text=True)
            queue = self.client.get('/').get_data(as_text=True)
        self.assertIn('Northwind Connectors', detail)
        self.assertIn('name="trade_show"', detail)
        self.assertIn('Space Tech Expo', detail)
        self.assertIn('N/A (Spreadsheet upload)', queue)


if __name__ == '__main__':
    unittest.main()
